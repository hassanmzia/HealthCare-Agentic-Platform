#!/usr/bin/env python3
"""Generate HealthCare Agentic Platform presentation."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# Colors
DARK_BLUE = RGBColor(0x02, 0x3E, 0x8A)
MEDIUM_BLUE = RGBColor(0x00, 0x77, 0xB6)
LIGHT_BLUE = RGBColor(0x90, 0xE0, 0xEF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT = RGBColor(0x1A, 0x1A, 0x2E)
GRAY_TEXT = RGBColor(0x55, 0x55, 0x55)
GREEN = RGBColor(0x2D, 0x6A, 0x4F)
GOLD = RGBColor(0xD4, 0xA0, 0x17)
RED_ACCENT = RGBColor(0xE7, 0x6F, 0x51)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_bg(slide, color=DARK_BLUE):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_text_box(slide, left, top, width, height, text, font_size=18, color=DARK_TEXT, bold=False, alignment=PP_ALIGN.LEFT, font_name='Calibri'):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox

def add_shape_box(slide, left, top, width, height, fill_color, border_color=None, text="", font_size=12, font_color=DARK_TEXT, bold=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    if text:
        tf = shape.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].text = text
        tf.paragraphs[0].font.size = Pt(font_size)
        tf.paragraphs[0].font.color.rgb = font_color
        tf.paragraphs[0].font.bold = bold
        tf.paragraphs[0].font.name = 'Calibri'
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return shape

# ============================================================
# SLIDE 1: Title Slide
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
add_bg(slide, DARK_BLUE)

# Gradient overlay bar
add_shape_box(slide, 0, 0, 13.333, 0.08, MEDIUM_BLUE)
add_shape_box(slide, 0, 7.42, 13.333, 0.08, MEDIUM_BLUE)

# Title
add_text_box(slide, 1, 1.5, 11, 1.2, "HealthCare Agentic Platform", font_size=48, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 1, 2.8, 11, 0.8, "AI-Powered Multi-Agent Clinical Decision Support System", font_size=24, color=LIGHT_BLUE, alignment=PP_ALIGN.CENTER)

# Badges
badges = ["FHIR R4 Compliant", "HIPAA Ready", "Multi-Agent AI", "Physician-in-the-Loop"]
for i, badge in enumerate(badges):
    x = 2.5 + i * 2.3
    add_shape_box(slide, x, 4.0, 2.0, 0.45, RGBColor(0x04, 0x50, 0x9A), MEDIUM_BLUE, badge, 11, WHITE, True)

add_text_box(slide, 1, 5.2, 11, 0.5, "Empowering clinicians with intelligent, evidence-based diagnostic\nand treatment recommendations", font_size=16, color=RGBColor(0xBB, 0xBB, 0xBB), alignment=PP_ALIGN.CENTER)
add_text_box(slide, 1, 6.3, 11, 0.4, "github.com/hassanmzia/HealthCare-Agentic-Platform", font_size=14, color=LIGHT_BLUE, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 2: The Problem
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_box(slide, 0, 0, 13.333, 1.0, DARK_BLUE)
add_text_box(slide, 0.5, 0.15, 12, 0.7, "The Challenge in Healthcare Today", font_size=32, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

problems = [
    ("Information Overload", "Clinicians spend 49% of time on\ndocumentation, not patient care", "49%"),
    ("Diagnostic Delays", "Average rare disease diagnosis\ntakes 5+ years", "5+ yrs"),
    ("Medication Errors", "7,000-9,000 annual deaths from\npreventable drug errors (US)", "9K"),
    ("Clinician Burnout", "63% of physicians report\nburnout symptoms", "63%"),
]

for i, (title, desc, stat) in enumerate(problems):
    x = 0.8 + i * 3.1
    add_shape_box(slide, x, 1.5, 2.8, 2.5, RGBColor(0xF8, 0xF9, 0xFA), MEDIUM_BLUE)
    add_text_box(slide, x + 0.2, 1.6, 2.4, 0.8, stat, font_size=36, color=RED_ACCENT, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.2, 2.3, 2.4, 0.4, title, font_size=16, color=DARK_BLUE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.1, 2.8, 2.6, 0.8, desc, font_size=11, color=GRAY_TEXT, alignment=PP_ALIGN.CENTER)

# Bottom stats
add_shape_box(slide, 0.5, 4.5, 12.3, 1.2, RGBColor(0xFE, 0xF0, 0xE0), RED_ACCENT)
add_text_box(slide, 1, 4.6, 11, 0.4, "The Cost of Inaction", font_size=18, color=RED_ACCENT, bold=True, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 1, 5.0, 11, 0.6, "$750B+ annual healthcare waste  |  12M diagnostic errors/year  |  30% clinical time spent searching for information", font_size=14, color=DARK_TEXT, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 3: Our Solution
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_box(slide, 0, 0, 13.333, 1.0, DARK_BLUE)
add_text_box(slide, 0.5, 0.15, 12, 0.7, "Our Solution: AI-Augmented Clinical Workflow", font_size=32, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

# Flow steps
flow_steps = [
    ("Patient Data\nIngestion", "Vitals, Labs,\nHistory, Imaging", RGBColor(0xD5, 0xE8, 0xD4)),
    ("Multi-Agent\nAI Analysis", "7+ Specialist\nAgents Collaborate", RGBColor(0xFF, 0xF2, 0xCC)),
    ("Comprehensive\nAssessment", "Diagnoses, Treatments,\nSafety Checks", RGBColor(0xF8, 0xCE, 0xCC)),
    ("Physician\nReview", "Approve, Modify,\nor Reject", RGBColor(0xDA, 0xE8, 0xFC)),
    ("Clinical\nOrders", "EHR Integration,\nFHIR Writeback", RGBColor(0xE1, 0xD5, 0xE7)),
]

for i, (title, desc, color) in enumerate(flow_steps):
    x = 0.6 + i * 2.5
    add_shape_box(slide, x, 1.5, 2.1, 1.8, color, MEDIUM_BLUE)
    add_text_box(slide, x + 0.1, 1.6, 1.9, 0.7, title, font_size=15, color=DARK_BLUE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.1, 2.3, 1.9, 0.7, desc, font_size=11, color=GRAY_TEXT, alignment=PP_ALIGN.CENTER)
    if i < 4:
        add_text_box(slide, x + 2.1, 1.9, 0.4, 0.5, "\u2192", font_size=28, color=MEDIUM_BLUE, bold=True, alignment=PP_ALIGN.CENTER)

# Key differentiators
diffs = [
    ("Multi-Agent AI", "7+ specialist agents collaborate\nlike a real clinical care team"),
    ("Physician-in-the-Loop", "Every recommendation requires\nphysician review before action"),
    ("FHIR-Native", "Built on HL7 FHIR R4 for\nseamless EHR interoperability"),
    ("Evidence-Based RAG", "Clinical guidelines embedded in\nvector DB for grounded recommendations"),
]

add_text_box(slide, 0.5, 3.8, 12, 0.5, "Key Differentiators", font_size=20, color=DARK_BLUE, bold=True)

for i, (title, desc) in enumerate(diffs):
    x = 0.5 + i * 3.1
    add_shape_box(slide, x, 4.4, 2.9, 1.3, RGBColor(0xF0, 0xF4, 0xFF), MEDIUM_BLUE)
    add_text_box(slide, x + 0.1, 4.5, 2.7, 0.4, title, font_size=14, color=DARK_BLUE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.1, 4.9, 2.7, 0.7, desc, font_size=10, color=GRAY_TEXT, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 4: AI Agents
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_box(slide, 0, 0, 13.333, 1.0, DARK_BLUE)
add_text_box(slide, 0.5, 0.15, 12, 0.7, "Multi-Agent Clinical AI Engine", font_size=32, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

# Supervisor
add_shape_box(slide, 4.5, 1.3, 4.3, 0.7, RGBColor(0xF8, 0xCE, 0xCC), RGBColor(0xB8, 0x54, 0x50), "SUPERVISOR AGENT\nOrchestrates All Specialists", 13, DARK_BLUE, True)

# Agent cards
agents = [
    ("Diagnostician", "Differential Dx\nConfidence Scoring"),
    ("Treatment", "Care Plans\nDosing Guidance"),
    ("Safety", "Drug Interactions\nAllergy Checks"),
    ("Medical Coding", "ICD-10 Codes\nCPT Codes"),
    ("Cardiology", "ECG Analysis\nCardiac Risk"),
    ("Radiology", "Imaging Review\nFinding Reports"),
    ("Pathology", "Lab Interpretation\nResult Trending"),
    ("Oncology", "Tumor Markers\nCancer Screening"),
    ("Gastroenterology", "GI Procedures\nGI Analysis"),
]

for i, (name, desc) in enumerate(agents):
    row = i // 5
    col = i % 5
    x = 0.8 + col * 2.4
    y = 2.5 + row * 1.6
    add_shape_box(slide, x, y, 2.1, 1.3, RGBColor(0xFF, 0xF9, 0xE6), GOLD)
    add_text_box(slide, x + 0.1, y + 0.15, 1.9, 0.4, name, font_size=13, color=DARK_BLUE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.1, y + 0.55, 1.9, 0.6, desc, font_size=10, color=GRAY_TEXT, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 5: Architecture
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_box(slide, 0, 0, 13.333, 1.0, DARK_BLUE)
add_text_box(slide, 0.5, 0.15, 12, 0.7, "System Architecture — 14 Microservices", font_size=32, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

# Architecture layers
layers = [
    ("PRESENTATION", "React 19 + TypeScript + Vite + Recharts", RGBColor(0xDA, 0xE8, 0xFC), 1.2,
     ["Doctor Portal", "Patient List", "Vitals Charts", "Labs Dashboard", "Alerts & Analytics"]),
    ("APPLICATION", "Django 5.0 + DRF + Celery + Redis", RGBColor(0xD5, 0xE8, 0xD4), 2.6,
     ["Patients API", "Clinical API", "Vitals API", "Labs API", "Medications API"]),
    ("AI ORCHESTRATION", "FastAPI + LangGraph + Claude/Ollama", RGBColor(0xFF, 0xF2, 0xCC), 4.0,
     ["Supervisor", "Diagnostician", "Treatment", "Safety", "Coding Agent"]),
    ("MCP TOOL SERVERS", "Model Context Protocol + HAPI FHIR R4", RGBColor(0xF5, 0xF5, 0xF5), 5.4,
     ["FHIR Server", "Labs Server", "RAG Server", "Pharmacy", "FHIR Adapter"]),
]

for label, tech, color, y, components in layers:
    add_shape_box(slide, 0.5, y, 12.3, 1.2, color, RGBColor(0x99, 0x99, 0x99))
    add_text_box(slide, 0.6, y + 0.05, 3, 0.3, label, font_size=11, color=DARK_BLUE, bold=True)
    add_text_box(slide, 0.6, y + 0.3, 4, 0.2, tech, font_size=8, color=GRAY_TEXT)
    for j, comp in enumerate(components):
        cx = 4.5 + j * 1.7
        add_shape_box(slide, cx, y + 0.25, 1.5, 0.65, WHITE, MEDIUM_BLUE, comp, 9, DARK_BLUE, True)

# Data stores on right
add_shape_box(slide, 10.5, 2.0, 2.2, 0.7, RGBColor(0xE1, 0xD5, 0xE7), RGBColor(0x96, 0x73, 0xA6), "PostgreSQL 16\n+ pgvector", 10, DARK_BLUE, True)
add_shape_box(slide, 10.5, 2.9, 2.2, 0.5, RGBColor(0xE1, 0xD5, 0xE7), RGBColor(0x96, 0x73, 0xA6), "Redis", 10, DARK_BLUE, True)
add_shape_box(slide, 10.5, 3.6, 2.2, 0.5, RGBColor(0xE1, 0xD5, 0xE7), RGBColor(0x96, 0x73, 0xA6), "ChromaDB", 10, DARK_BLUE, True)

# ============================================================
# SLIDE 6: Technology Stack
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_box(slide, 0, 0, 13.333, 1.0, DARK_BLUE)
add_text_box(slide, 0.5, 0.15, 12, 0.7, "Technology Stack", font_size=32, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

tech_categories = [
    ("Frontend", RGBColor(0xDA, 0xE8, 0xFC), [
        "React 19 — UI framework",
        "TypeScript 5.9 — Type safety",
        "Vite 7 — Build tooling",
        "TanStack Query — Server state",
        "Recharts — Data visualization",
        "Nginx — Reverse proxy + SSL",
    ]),
    ("Backend", RGBColor(0xD5, 0xE8, 0xD4), [
        "Django 5.0 — REST API",
        "Django REST Framework — Serialization",
        "PostgreSQL 16 — Primary database",
        "pgvector — Vector embeddings",
        "Celery + Redis — Async tasks",
        "JWT — Authentication",
    ]),
    ("AI / ML", RGBColor(0xFF, 0xF2, 0xCC), [
        "Claude (Anthropic) — Primary LLM",
        "Ollama — Local LLM alternative",
        "LangGraph — Agent orchestration",
        "ChromaDB — Vector database",
        "Sentence-Transformers — Embeddings",
        "Model Context Protocol — Tool comm",
    ]),
    ("Infrastructure", RGBColor(0xE1, 0xD5, 0xE7), [
        "Docker Compose — Orchestration",
        "FastAPI — AI orchestrator API",
        "HAPI FHIR — FHIR R4 server",
        "SSL/TLS — Encryption",
        "Healthchecks — Service monitoring",
        "Docker Network — Service mesh",
    ]),
]

for i, (cat, color, items) in enumerate(tech_categories):
    x = 0.5 + i * 3.15
    add_shape_box(slide, x, 1.3, 2.9, 4.8, color, MEDIUM_BLUE)
    add_text_box(slide, x + 0.1, 1.4, 2.7, 0.4, cat, font_size=18, color=DARK_BLUE, bold=True, alignment=PP_ALIGN.CENTER)
    for j, item in enumerate(items):
        add_text_box(slide, x + 0.2, 1.9 + j * 0.55, 2.5, 0.5, item, font_size=11, color=DARK_TEXT)

# ============================================================
# SLIDE 7: Features & Capabilities
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_box(slide, 0, 0, 13.333, 1.0, DARK_BLUE)
add_text_box(slide, 0.5, 0.15, 12, 0.7, "Platform Capabilities", font_size=32, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

capabilities = [
    ("Clinical Intelligence", GREEN, [
        "Multi-agent differential diagnosis",
        "Confidence-scored recommendations",
        "Evidence-based treatment planning",
        "Automated medical coding (ICD-10/CPT)",
        "Clinical guidelines RAG",
    ]),
    ("Real-Time Monitoring", MEDIUM_BLUE, [
        "IoT device integration",
        "Continuous vital sign tracking",
        "Configurable alert thresholds",
        "Critical value detection",
        "Automated notification system",
    ]),
    ("Physician Workflow", GOLD, [
        "Review & approve AI recommendations",
        "Auto-generated clinical documentation",
        "SOAP notes & discharge summaries",
        "Digital physician signatures",
        "Complete audit trail",
    ]),
    ("Security & Compliance", RED_ACCENT, [
        "HIPAA-ready architecture",
        "JWT + role-based access (6 roles)",
        "Patient data segregation",
        "SSL/TLS encryption",
        "Before/after audit logging",
    ]),
]

for i, (title, color, items) in enumerate(capabilities):
    x = 0.5 + i * 3.15
    add_shape_box(slide, x, 1.3, 2.9, 0.6, color, text=title, font_size=15, font_color=WHITE, bold=True)
    for j, item in enumerate(items):
        y = 2.1 + j * 0.45
        add_text_box(slide, x + 0.3, y, 2.5, 0.4, "\u2713  " + item, font_size=11, color=DARK_TEXT)

# ROI section
add_text_box(slide, 0.5, 4.6, 12, 0.5, "Expected ROI Metrics", font_size=18, color=DARK_BLUE, bold=True, alignment=PP_ALIGN.CENTER)

roi_items = [
    ("-50%", "Documentation\nTime"),
    ("+40%", "Diagnostic\nAccuracy"),
    ("+20%", "Coding Revenue\nRecovery"),
    ("95%+", "Drug Interaction\nDetection"),
    ("+30%", "Clinician\nSatisfaction"),
]

for i, (stat, label) in enumerate(roi_items):
    x = 1.0 + i * 2.4
    add_shape_box(slide, x, 5.2, 2.0, 1.2, RGBColor(0xE8, 0xF4, 0xF8), MEDIUM_BLUE)
    add_text_box(slide, x + 0.1, 5.25, 1.8, 0.5, stat, font_size=24, color=MEDIUM_BLUE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.1, 5.75, 1.8, 0.5, label, font_size=10, color=GRAY_TEXT, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 8: Service Ports & Deployment
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_box(slide, 0, 0, 13.333, 1.0, DARK_BLUE)
add_text_box(slide, 0.5, 0.15, 12, 0.7, "Deployment & Service Architecture", font_size=32, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

services = [
    ("clinician-ui", ":3030", "React + Nginx", RGBColor(0xDA, 0xE8, 0xFC)),
    ("backend", ":8000", "Django 5.0", RGBColor(0xD5, 0xE8, 0xD4)),
    ("worker", "—", "Celery", RGBColor(0xD5, 0xE8, 0xD4)),
    ("orchestrator", ":8003", "FastAPI + LangGraph", RGBColor(0xFF, 0xF2, 0xCC)),
    ("iot-simulator", ":8004", "FastAPI", RGBColor(0xE6, 0xFF, 0xCC)),
    ("mcp-fhir-server", ":8005", "FHIR Read", RGBColor(0xF5, 0xF5, 0xF5)),
    ("mcp-labs-server", ":8006", "Lab Results", RGBColor(0xF5, 0xF5, 0xF5)),
    ("mcp-rag-server", ":8007", "Clinical RAG", RGBColor(0xF5, 0xF5, 0xF5)),
    ("mcp-pharmacy", ":8008", "Drug Formulary", RGBColor(0xF5, 0xF5, 0xF5)),
    ("mcp-fhir-adapter", ":8002", "FHIR Writeback", RGBColor(0xF5, 0xF5, 0xF5)),
    ("hapi-fhir", ":18090", "FHIR R4 Server", RGBColor(0xFF, 0xE6, 0xCC)),
    ("db", ":5432", "PostgreSQL 16", RGBColor(0xE1, 0xD5, 0xE7)),
    ("redis", ":6379", "Message Broker", RGBColor(0xE1, 0xD5, 0xE7)),
    ("chromadb", ":8009", "Vector DB", RGBColor(0xE1, 0xD5, 0xE7)),
]

for i, (name, port, desc, color) in enumerate(services):
    row = i // 5
    col = i % 5
    x = 0.5 + col * 2.55
    y = 1.3 + row * 1.4
    add_shape_box(slide, x, y, 2.3, 1.1, color, MEDIUM_BLUE)
    add_text_box(slide, x + 0.1, y + 0.1, 2.1, 0.3, name, font_size=12, color=DARK_BLUE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.1, y + 0.4, 2.1, 0.3, port, font_size=14, color=MEDIUM_BLUE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.1, y + 0.7, 2.1, 0.3, desc, font_size=9, color=GRAY_TEXT, alignment=PP_ALIGN.CENTER)

# Quick start
add_shape_box(slide, 0.5, 5.6, 12.3, 1.5, RGBColor(0x1A, 0x1A, 0x2E))
add_text_box(slide, 1, 5.7, 5, 0.4, "Quick Start", font_size=16, color=LIGHT_BLUE, bold=True)
add_text_box(slide, 1, 6.1, 11, 0.9,
    "$ git clone https://github.com/hassanmzia/HealthCare-Agentic-Platform.git\n"
    "$ cd HealthCare-Agentic-Platform\n"
    "$ docker-compose up -d                    # Launch all 14 services\n"
    "# Access dashboard at https://localhost:3030",
    font_size=12, color=LIGHT_BLUE, font_name='Courier New')

# ============================================================
# SLIDE 9: Closing / Thank You
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BLUE)
add_shape_box(slide, 0, 0, 13.333, 0.08, MEDIUM_BLUE)
add_shape_box(slide, 0, 7.42, 13.333, 0.08, MEDIUM_BLUE)

add_text_box(slide, 1, 1.5, 11, 1.0, "HealthCare Agentic Platform", font_size=44, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 1, 2.6, 11, 0.6, "AI-Powered Clinical Intelligence.\nPhysician-Approved Care.", font_size=24, color=LIGHT_BLUE, alignment=PP_ALIGN.CENTER)

# Summary stats
summary = [
    ("7+", "AI Specialist\nAgents"),
    ("14", "Docker\nMicroservices"),
    ("FHIR R4", "Healthcare\nInteroperability"),
    ("HIPAA", "Compliance\nReady"),
]

for i, (num, label) in enumerate(summary):
    x = 2.0 + i * 2.5
    add_shape_box(slide, x, 3.8, 2.1, 1.5, RGBColor(0x04, 0x50, 0x9A), MEDIUM_BLUE)
    add_text_box(slide, x + 0.1, 3.9, 1.9, 0.7, num, font_size=30, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.1, 4.5, 1.9, 0.6, label, font_size=12, color=LIGHT_BLUE, alignment=PP_ALIGN.CENTER)

add_text_box(slide, 1, 5.8, 11, 0.4, "github.com/hassanmzia/HealthCare-Agentic-Platform", font_size=16, color=LIGHT_BLUE, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 1, 6.4, 11, 0.4, "Built with Django  |  React  |  LangGraph  |  Claude AI  |  FHIR R4", font_size=14, color=RGBColor(0x88, 0x88, 0x88), alignment=PP_ALIGN.CENTER)

# Save
output_path = os.path.join(os.path.dirname(__file__), "HealthCare_Agentic_Platform.pptx")
prs.save(output_path)
print(f"Presentation saved to: {output_path}")
